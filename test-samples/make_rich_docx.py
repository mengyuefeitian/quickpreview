# -*- coding: utf-8 -*-
"""生成一个排版要素齐全的 Word 样本，用于对比预览还原度与耗时。"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = os.path.dirname(os.path.abspath(__file__))

doc = Document()

# 页面：A4
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)

# 标题
h = doc.add_heading('季度业务复盘报告', level=0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 正文段落：不同字号 + 颜色
p = doc.add_paragraph()
r = p.add_run('本报告汇总了 2026 年第三季度各业务线的核心指标，')
r.font.size = Pt(14)
r = p.add_run('重点标注了异常波动项')
r.font.size = Pt(14)
r.bold = True
r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
r = p.add_run('。')
r.font.size = Pt(14)

# 斜体 + 下划线
p2 = doc.add_paragraph()
r2 = p2.add_run('斜体说明文字：')
r2.italic = True
r2.font.size = Pt(12)
r2b = p2.add_run('下划线强调')
r2b.underline = True
r2b.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

# 小标题
doc.add_heading('一、核心指标', level=1)
doc.add_heading('1.1 收入构成', level=2)

# 项目符号列表
for t in ['订阅收入：占比 62%，同比 +18%', '广告收入：占比 25%，同比 -4%', '增值服务：占比 13%，同比 +31%']:
    doc.add_paragraph(t, style='List Bullet')

# 编号列表
for t in ['确认口径调整后的可比基数', '剔除一次性减值影响', '按区域重新分摊公共成本']:
    doc.add_paragraph(t, style='List Number')

# 缩进段落
ind = doc.add_paragraph('这是一段带左右缩进与首行缩进的正文，用于验证缩进还原效果。')
ind.paragraph_format.left_indent = Cm(2)
ind.paragraph_format.right_indent = Cm(1.5)
ind.paragraph_format.first_line_indent = Cm(1)
ind.paragraph_format.space_after = Pt(12)

# 右对齐
ra = doc.add_paragraph('—— 数据截止 2026-09-30')
ra.alignment = WD_ALIGN_PARAGRAPH.RIGHT

# 表格
doc.add_heading('二、分区域明细', level=1)
tbl = doc.add_table(rows=1, cols=4)
tbl.style = 'Table Grid'
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tbl.rows[0].cells
for i, t in enumerate(['区域', '收入(万元)', '同比', '达成率']):
    hdr[i].text = t
    for para in hdr[i].paragraphs:
        for run in para.runs:
            run.bold = True
rows = [
    ['华东', '4,820', '+12%', '104%'],
    ['华南', '3,150', '+8%', '97%'],
    ['华北', '2,760', '-3%', '88%'],
    ['西部', '1,430', '+25%', '112%'],
]
for row in rows:
    cells = tbl.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v

# 图片
doc.add_heading('三、趋势图', level=1)
img_path = os.path.join(HERE, 'sample.png')
if os.path.exists(img_path):
    doc.add_picture(img_path, width=Cm(12))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# 大段正文（用于测速）
doc.add_heading('四、附录：明细说明', level=1)
for i in range(60):
    pi = doc.add_paragraph(
        f'附录第 {i+1} 条：本条目用于增加文档体量，验证解析耗时。'
        f'包含若干中文与 English 混排文本，以及数字 123456 与符号 。'
    )
    pi.paragraph_format.space_after = Pt(4)

out = os.path.join(HERE, 'rich.docx')
doc.save(out)
print('生成', out, os.path.getsize(out), '字节')
